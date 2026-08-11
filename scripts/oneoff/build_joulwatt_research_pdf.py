#!/usr/bin/env python3
"""Build the latest Joulwatt research report PDF."""

from __future__ import annotations

import os

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "output/reports/joulwatt_research_report_pdf_20260708.md"
DOCX = ROOT / "output/reports/joulwatt_research_report_pdf_20260708.docx"
PDF_DIR = ROOT / "output/pdf"
PDF = PDF_DIR / "joulwatt_research_report_pdf_20260708.pdf"

FONT = "Arial"
FONT_EAST = "PingFang SC"
BLUE = "183B59"
TEAL = "0F766E"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F8"
BORDER = "B8C2CC"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, width_in: float) -> None:
    cell.width = Inches(width_in)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_in * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color=None) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_marked_text(paragraph, text: str, size=10.2, color=None) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        clean = part[2:-2] if bold else part[1:-1] if code else part
        run = paragraph.add_run(clean)
        set_run_font(run, size=size, bold=bold, color=color)
        if code:
            run.font.name = "Menlo"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST)
        if "🔼" in clean:
            run.font.color.rgb = RGBColor.from_string("15803D")
        if "🔽" in clean:
            run.font.color.rgb = RGBColor.from_string("B91C1C")


def clean_text(text: str) -> str:
    return text.strip().replace("<br>", " ")


def style_paragraph(paragraph, after=5, before=0, line=1.06) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(after)
    fmt.space_before = Pt(before)
    fmt.line_spacing = line


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph()
    if level == 1:
        style_paragraph(p, before=8, after=5)
        run = p.add_run(text)
        set_run_font(run, 14.5, True, BLUE)
    else:
        style_paragraph(p, before=6, after=3)
        run = p.add_run(text)
        set_run_font(run, 11.5, True, TEAL)


def column_widths(headers: list[str]) -> list[float]:
    n = len(headers)
    total = 9.85
    if n == 5:
        return [1.15, 2.05, 2.05, 2.0, 2.6]
    if n == 4:
        return [1.55, 1.5, 3.2, 3.6]
    if n == 3:
        return [1.6, 4.1, 4.15]
    return [total / n] * n


def add_table(doc: Document, lines: list[str]) -> None:
    parsed = [[clean_text(c) for c in line.strip().strip("|").split("|")] for line in lines]
    if len(parsed) < 2:
        return
    headers = parsed[0]
    rows = [r for r in parsed[2:] if any(c.strip() for c in r)]
    widths = column_widths(headers)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_width(cell, widths[idx])
        set_cell_margins(cell)
        set_cell_shading(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_marked_text(p, text, size=8.4)
        for run in p.runs:
            run.bold = True

    for ridx, row_values in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row_values[: len(headers)]):
            cell = cells[idx]
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell, top=65, bottom=65)
            if ridx in (2, 3, 4) and idx == 3:
                set_cell_shading(cell, "FFF7ED")
            elif ridx % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 1, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
            add_marked_text(p, text, size=7.7 if len(headers) >= 5 else 8.4)
    doc.add_paragraph()


def build() -> None:
    PDF_DIR.mkdir(parents=True, exist_ok=True)
    DOCX.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST)
    normal.font.size = Pt(10.2)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.06

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_marked_text(footer, "杰华特业绩预测研报 | 2026-07-08", size=8, color="666666")

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line == "---":
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            style_paragraph(p, after=2)
            run = p.add_run(line[2:].strip())
            set_run_font(run, 21, True, "0B2545")
            i += 1
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:].strip(), 1)
            i += 1
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:].strip(), 2)
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        if line.startswith("- "):
            p = doc.add_paragraph()
            style_paragraph(p, after=3)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.12)
            add_marked_text(p, "• " + line[2:].strip(), size=10)
            i += 1
            continue

        p = doc.add_paragraph()
        style_paragraph(p)
        add_marked_text(p, clean_text(line), size=10.2)
        i += 1

    doc.save(DOCX)

    subprocess.run(
        [
            os.environ.get("SOFFICE_BIN", "/Users/zyb/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/soffice"),
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(PDF_DIR),
            str(DOCX),
        ],
        check=True,
        env={"TMPDIR": "/private/tmp", "HOME": str(ROOT / "tmp/lo-home")},
    )
    generated = PDF_DIR / f"{DOCX.stem}.pdf"
    if generated != PDF:
        generated.replace(PDF)


if __name__ == "__main__":
    build()
    print(DOCX)
    print(PDF)
