from pathlib import Path
from math import ceil

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


import os
# 一次性交付脚本也遵循统一运行目录；未配置时兼容旧 output/tmp。
PROJECT_ROOT = Path(__file__).resolve().parents[2]
RUNTIME_ROOT = os.environ.get("IRA_RUNTIME_ROOT")
OUT = Path(os.environ.get("IRA_ARTIFACT_ROOT", Path(RUNTIME_ROOT) / "artifacts" if RUNTIME_ROOT else PROJECT_ROOT / "output")) / "pdf"
TMP = Path(os.environ.get("IRA_CACHE_ROOT", Path(RUNTIME_ROOT) / "cache" if RUNTIME_ROOT else PROJECT_ROOT / "tmp")) / "amd_report"
OUT.mkdir(parents=True, exist_ok=True)
TMP.mkdir(parents=True, exist_ok=True)

DOCX_PATH = OUT / "AMD_2026Q2_20260804_Codex独立研报.docx"

NAVY = "15344A"
BLUE = "1F6A8A"
TEAL = "2A9D8F"
GOLD = "D69E2E"
RED = "B23A48"
INK = "17232D"
MID = "536270"
LIGHT = "EAF0F4"
PALE = "F5F7F9"
WHITE = "FFFFFF"
GRID = "CAD4DC"

FONT_CN = "STHeiti"
FONT_LATIN = "STHeiti"
CONTENT_DXA = 9360


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths):
    assert sum(widths) == CONTENT_DXA
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[i]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[i] / 1440)
            set_cell_margins(cell)


def set_table_borders(table, color=GRID, size=4):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_run_font(run, size=10.5, bold=False, color=INK, italic=False, latin=FONT_LATIN):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hint"), "eastAsia")
    lang = run._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.get_or_add_rPr().append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_para(p, before=0, after=6, line=1.10, align=None, keep_next=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    if align is not None:
        p.alignment = align


def add_text(doc, text, size=10.5, color=INK, bold=False, italic=False,
             before=0, after=6, align=None, keep_next=False):
    p = doc.add_paragraph()
    set_para(p, before, after, 1.10, align, keep_next)
    r = p.add_run(text)
    set_run_font(r, size, bold, color, italic)
    return p


def add_rich_para(doc, segments, before=0, after=6, align=None, keep_next=False):
    p = doc.add_paragraph()
    set_para(p, before, after, 1.10, align, keep_next)
    for text, kwargs in segments:
        r = p.add_run(text)
        set_run_font(r, **kwargs)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    sizes = {1: 16, 2: 13, 3: 11.5}
    colors = {1: NAVY, 2: BLUE, 3: NAVY}
    set_run_font(r, sizes[level], True, colors[level])
    return p


def add_source(doc, text):
    return add_text(doc, text, size=8.2, color=MID, italic=True, before=2, after=7)


def add_callout(doc, label, text, fill=LIGHT, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [CONTENT_DXA])
    set_table_borders(table, color=fill, size=2)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    set_para(p, after=0, line=1.10)
    r = p.add_run(label + "  ")
    set_run_font(r, 10.5, True, accent)
    r = p.add_run(text)
    set_run_font(r, 10.5, False, INK)
    add_text(doc, "", size=2, after=2)
    return table


def add_table(doc, headers, rows, widths, aligns=None, font_size=8.8, zebra=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths)
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j, h in enumerate(headers):
        cell = hdr.cells[j]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        set_para(p, after=0, line=1.0)
        r = p.add_run(str(h))
        set_run_font(r, font_size, True, WHITE)
    for i, row_data in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row_data):
            cell = cells[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if zebra and i % 2 == 1:
                set_cell_shading(cell, PALE)
            p = cell.paragraphs[0]
            alignment = aligns[j] if aligns else (WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            p.alignment = alignment
            set_para(p, after=0, line=1.0)
            r = p.add_run(str(value))
            set_run_font(r, font_size, j == 0, INK)
    return table


def add_page_break(doc):
    doc.add_page_break()


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, 8.5, False, MID)


def make_bar_chart(path, labels, series, title, subtitle=None):
    width, height = 1400, 720
    hexcolor = lambda c: c if str(c).startswith("#") else f"#{c}"
    img = Image.new("RGB", (width, height), hexcolor(WHITE))
    draw = ImageDraw.Draw(img)
    font_path = "/System/Library/Fonts/STHeiti Light.ttc"
    try:
        title_font = ImageFont.truetype(font_path, 44)
        sub_font = ImageFont.truetype(font_path, 24)
        label_font = ImageFont.truetype(font_path, 23)
        value_font = ImageFont.truetype(font_path, 22)
    except OSError:
        title_font = sub_font = label_font = value_font = ImageFont.load_default()
    draw.text((70, 38), title, fill=hexcolor(NAVY), font=title_font)
    if subtitle:
        draw.text((72, 98), subtitle, fill=hexcolor(MID), font=sub_font)
    left, top, right, bottom = 110, 170, 1320, 620
    draw.line((left, bottom, right, bottom), fill=hexcolor(GRID), width=3)
    max_v = max(max(vals) for _, vals, _ in series)
    group_w = (right - left) / len(labels)
    bar_w = group_w / (len(series) + 1.7)
    for i, lab in enumerate(labels):
        x_center = left + group_w * (i + 0.5)
        draw.text((x_center - 35, bottom + 20), lab, fill=hexcolor(INK), font=label_font)
        for sidx, (name, vals, color) in enumerate(series):
            v = vals[i]
            x0 = x_center - (len(series) * bar_w) / 2 + sidx * bar_w
            x1 = x0 + bar_w * 0.78
            y0 = bottom - (v / max_v) * (bottom - top)
            draw.rounded_rectangle((x0, y0, x1, bottom), radius=9, fill=hexcolor(color))
            text = f"{v:g}"
            draw.text((x0 + 2, y0 - 31), text, fill=hexcolor(INK), font=value_font)
    lx = 820
    for idx, (name, _, color) in enumerate(series):
        y = 52 + idx * 38
        draw.rounded_rectangle((lx, y, lx + 28, y + 20), radius=4, fill=hexcolor(color))
        draw.text((lx + 40, y - 5), name, fill=hexcolor(INK), font=sub_font)
    img.save(path)


def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.82)
    sec.bottom_margin = Inches(0.78)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, NAVY, 8, 4),
    ):
        st = styles[name]
        st.font.name = FONT_LATIN
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        st._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para(hp, after=0, line=1.0)
    r = hp.add_run("AMD DEEP RESEARCH  |  2026.08.04")
    set_run_font(r, 8.5, True, MID)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_para(fp, after=0, line=1.0)
    r = fp.add_run("独立研究草案  ·  仅供研究讨论  |  ")
    set_run_font(r, 8.2, False, MID)
    add_page_field(fp)
    return doc


def build():
    chart_revenue = TMP / "revenue_scenarios.png"
    chart_value = TMP / "valuation_scenarios.png"
    make_bar_chart(
        chart_revenue,
        ["2025A", "2026E", "2027E", "2028E"],
        [
            ("审慎", [34.6, 46, 58, 72], "AAB7C4"),
            ("基准", [34.6, 50.5, 82, 110], TEAL),
            ("乐观", [34.6, 60, 110, 145], GOLD),
        ],
        "AMD收入三情景（十亿美元）",
        "基准情景以Q1实际、Q2指引和可验证的MI450/EPYC爬坡为锚",
    )
    make_bar_chart(
        chart_value,
        ["审慎", "基准", "乐观"],
        [
            ("2027E EPS", [9.0, 14.1, 18.0], BLUE),
            ("目标价值/10", [25.2, 46.5, 64.8], GOLD),
        ],
        "估值敏感性：盈利兑现比倍数更重要",
        "目标价值按2027E EPS×28x/33x/36x；图中目标价值按10美元缩放",
    )

    doc = setup_document()

    # Cover
    add_text(doc, "CODEX版 · 独立公司研究", size=10, color=GOLD, bold=True, before=40, after=22,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "超威半导体（AMD）", size=29, color=NAVY, bold=True, after=6,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "AI算力第二供应链的真实弹性与兑现门槛", size=16, color=BLUE, bold=True,
             after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "从GW合作公告回到出货、收入确认、毛利率与摊薄后EPS", size=11.5,
             color=MID, after=52, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_callout(doc, "投资结论", "中性 / 观察。基准合理价值465美元，合理区间430-530美元；当前价格484.64美元，市场已计入相当部分2027年MI450与EPYC增长。", fill=LIGHT, accent=BLUE)
    add_text(doc, "基准日：2026年8月4日（AMD 2026Q2财报发布前）", size=10, color=MID,
             before=24, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "当前价格：484.64美元  |  市值：约7,997亿美元  |  未来12个月核心事件：MI450量产与Helios机架交付",
             size=9.5, color=MID, after=60, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "核心观点", size=12, color=NAVY, bold=True, after=8)
    cover_points = [
        ("产业位置：", "AMD已从CPU/GPU芯片商转向CPU、GPU、网络、机架和ROCm软件的系统供应商。"),
        ("增长主线：", "2026H2-2027增长主要来自MI450/Helios和EPYC服务器CPU，而非客户端业务。"),
        ("主要分歧：", "GW协议代表多年期潜在需求，不等于当期确认收入；权证、战略投资与客户集中降低每股价值确定性。"),
        ("估值判断：", "现价并不便宜。只有在2027年收入接近1000亿美元、摊薄后EPS接近18美元时，600美元以上目标才成立。"),
    ]
    for label, text in cover_points:
        add_rich_para(doc, [(label, {"size": 10.3, "bold": True, "color": BLUE}),
                            (text, {"size": 10.3, "bold": False, "color": INK})], after=6)
    add_source(doc, "资料来源：AMD公司公告、AMD 2026Q1财务报告、AMD Advancing AI 2026、UBS（2026-06-24）、公开市场数据；本报告测算。")

    add_page_break(doc)

    # Executive summary
    add_heading(doc, "1. 投资摘要", 1)
    add_callout(doc, "一句话结论", "AMD的第二供应链地位正在得到客户合同验证，但当前股价交易的是2027年而不是2026年；应把MI450交付和权证摊薄同时纳入估值。", fill="E8F3F1", accent=TEAL)
    add_heading(doc, "关键指标", 2)
    add_table(doc,
              ["指标", "2025A", "2026E", "2027E", "2028E"],
              [
                  ["收入（亿美元）", "346", "505", "820", "1,100"],
                  ["同比", "34%", "46%", "62%", "34%"],
                  ["数据中心收入（亿美元）", "166", "320", "645", "903"],
                  ["Non-GAAP毛利率", "52.4%", "55.8%", "55.5%", "56.2%"],
                  ["摊薄后Non-GAAP EPS（美元）", "4.17", "7.6", "14.1", "20.2"],
                  ["估算摊薄股数（亿股）", "16.4", "16.6", "17.2", "18.0"],
              ], [3000, 1590, 1590, 1590, 1590], font_size=9.1)
    add_source(doc, "注：摊薄股数为研究情景，不代表权证一定全部归属；2025A EPS采用可比口径。")

    add_heading(doc, "投资逻辑的三层证据", 2)
    for title, body in [
        ("第一层｜CPU份额与ASP", "EPYC持续受益于核心数、能效与x86软件兼容性。Agentic AI提高控制面、检索、数据处理和传统软件负载，CPU需求有独立增量，而不仅是GPU附属。"),
        ("第二层｜GPU与机架", "OpenAI、Meta、Anthropic及微软的合作把MI450从产品路线图推进到客户部署阶段。真正决定收入的是机架物料价值、交付节奏和客户验收。"),
        ("第三层｜开放软件", "ROCm与Coding Agent可以降低迁移成本，但短期不能等同于CUDA生态被替代。生产稳定性、算子覆盖、多机通信和开发工具仍需持续验证。"),
    ]:
        add_rich_para(doc, [(title, {"size": 10.4, "bold": True, "color": BLUE}),
                            ("  " + body, {"size": 10.4, "bold": False, "color": INK})], after=8)

    add_heading(doc, "最关键的跟踪变量", 2)
    add_table(doc,
              ["变量", "基准假设", "验证信号", "失效信号"],
              [
                  ["MI450/Helios", "2026H2开始贡献，2027放量", "单季数据中心收入连续上台阶", "出货有但验收/确认延后"],
                  ["服务器CPU", "2026收入约160亿美元", "单位与ASP同步提升", "Arm/定制芯片侵蚀独立CPU负载"],
                  ["ROCm", "迁移效率改善但仍需工程投入", "生产客户与利用率披露增加", "仅演示适配、缺少规模生产"],
                  ["股本摊薄", "权证分阶段归属", "收入与EPS增量覆盖摊薄", "股数增长快于利润兑现"],
              ], [1750, 2450, 2580, 2580], font_size=8.5)

    add_page_break(doc)

    # Industry and platform
    add_heading(doc, "2. 公司定位：从芯片组合走向异构系统", 1)
    add_text(doc, "AMD的竞争方式不是复制英伟达的单一GPU优势，而是用EPYC CPU、Instinct GPU、Pensando NIC/DPU、Helios机架与ROCm软件形成可组合的平台。该组合对希望降低单一供应商依赖的云厂商尤其有吸引力。")
    add_table(doc,
              ["层级", "核心产品", "客户价值", "核心验证点"],
              [
                  ["计算控制面", "EPYC Turin/Venice/Verano", "任务编排、传统软件、数据处理", "单位、ASP、云实例份额"],
                  ["加速计算", "MI350/MI450/MI500", "训练与高吞吐推理", "性能/美元、供货与利用率"],
                  ["网络与数据面", "Pensando NIC/DPU", "机架级互联、卸载与安全", "互联规模、客户采用"],
                  ["系统", "Helios", "缩短部署周期、统一交付", "BOM、毛利率、验收节奏"],
                  ["软件", "ROCm.AI", "开放生态与迁移自动化", "算子覆盖、稳定性、开发工具"],
              ], [1700, 2200, 2820, 2640], font_size=8.7)
    add_source(doc, "资料来源：AMD Advancing AI 2026；本报告整理。")

    add_heading(doc, "产品节奏", 2)
    add_table(doc,
              ["时间", "GPU/系统", "CPU/网络", "研究判断"],
              [
                  ["2026H2", "MI450、Helios首批出货", "Venice、Pensando", "收入起点，重点看验收而非公告GW"],
                  ["2027", "MI500、Helios 500", "Verano、Como/Monza", "规模放量与平台毛利率关键年"],
                  ["2028", "MI600、Helios 600", "Florence/Ferrara、Palma/Levanzo", "验证年度迭代与客户留存"],
              ], [1450, 2300, 2500, 3110], font_size=8.8)

    add_heading(doc, "TAM可以扩大，但不能替代份额模型", 2)
    add_text(doc, "AMD给出的2030年AI芯片TAM约1.4万亿美元、数据中心CPU TAM约2200亿美元，说明行业空间足够大；但TAM不是收入预测。研究模型必须回答AMD在每一代产品中的单位份额、ASP、机架价值量和收入确认时点。")
    add_callout(doc, "判断", "“市场足够大”只解决天花板问题；“客户何时验收、AMD确认多少收入、付出多少稀释和投资”才决定每股价值。", fill="FFF7E6", accent=GOLD)

    add_page_break(doc)

    # Customer commitments
    add_heading(doc, "3. 客户合作：GW是需求框架，不是收入数字", 1)
    add_table(doc,
              ["客户", "合作规模", "首批节奏", "对价/条件", "模型处理"],
              [
                  ["OpenAI", "多年期6GW", "首个1GW自2026H2开始", "最高1.6亿股权证；技术、采购及股价条件", "2026仅计部分出货，2027-30分期"],
                  ["Meta", "最高6GW", "首个1GW相关出货自2026H2", "最高1.6亿股权证；里程碑归属", "定制MI450，分期计入"],
                  ["Anthropic", "最高2GW", "首个1GW自2027H1开始", "AMD未来最高50亿美元股权投资", "收入与投资现金流分开"],
                  ["Microsoft", "计划部署Helios", "未披露明确数量", "合作与产品验证", "不提前计入大额承诺"],
                  ["Oracle", "首期5万颗MI450", "2026Q3开始", "云服务部署", "按季度供货爬坡"],
              ], [1280, 1400, 1900, 2720, 2060], font_size=8.1)
    add_source(doc, "资料来源：AMD关于OpenAI、Meta、Anthropic、Microsoft和Oracle的官方公告；规模均为潜在/计划部署口径。")

    add_heading(doc, "为什么不能用“美元/GW”直接乘", 2)
    add_text(doc, "一个GW可对应不同GPU数量、功耗配置、网络方案与系统边界。若AMD只确认GPU和部分网络收入，收入价值量较低而毛利率较高；若确认完整机架收入，收入更高但包含更多低毛利率部件。缺少合同结构时，收入和毛利率必须成对假设。")
    add_table(doc,
              ["确认边界", "收入弹性", "毛利率倾向", "主要风险"],
              [
                  ["芯片为主", "较低", "较高", "GW标题难直接映射为AMD收入"],
                  ["GPU+网络", "中等", "中高", "互联与客户定制影响价值量"],
                  ["完整机架/系统", "较高", "可能较低", "供应链、验收及营运资金压力更大"],
              ], [1900, 1700, 1900, 3860], font_size=8.8)

    add_heading(doc, "客户集中与信用链条", 2)
    add_text(doc, "OpenAI与Meta的潜在规模足以显著改变AMD收入结构，也会提高单一客户、融资和项目节奏风险。AMD对NeoCloud伙伴的投资或担保有助于需求形成，但同时把供应商风险扩展为投资、信用和回收风险。")

    add_page_break(doc)

    # Financial model
    add_heading(doc, "4. 财务模型：用季度桥约束全年预测", 1)
    doc.add_picture(str(chart_revenue), width=Inches(6.45))
    add_source(doc, "资料来源：AMD、UBS；2026E-2028E为本报告情景测算。")

    add_heading(doc, "2026收入桥", 2)
    add_table(doc,
              ["项目", "收入（亿美元）", "状态/含义"],
              [
                  ["2026Q1", "102.5", "已公布"],
                  ["2026Q2", "112.0", "公司指引中值"],
                  ["基准H2", "290.5", "单季均值145.3，较H1单季均值高35%"],
                  ["基准全年", "505.0", "接近近期可验证卖方底模"],
                  ["原华泰全年", "638.9", "要求H2收入424.4，单季均值212.2"],
              ], [2600, 2100, 4660], font_size=9.0)
    add_callout(doc, "关键差异", "本报告没有把首批GW“开始出货”解释为当年完整部署；2026收入主要由现有MI350、EPYC增长和MI450首批贡献构成。", fill="E8F3F1", accent=TEAL)

    add_heading(doc, "分部收入基准情景", 2)
    add_table(doc,
              ["亿美元", "2025A", "2026E", "2027E", "2028E", "核心驱动"],
              [
                  ["数据中心", "166", "320", "645", "903", "MI450/Helios、EPYC、Pensando"],
                  ["客户端", "106", "118", "105", "115", "份额提升后趋稳，PC周期正常化"],
                  ["游戏", "39", "27", "23", "32", "半定制下行后恢复"],
                  ["嵌入式", "35", "40", "47", "50", "库存周期修复"],
                  ["合计", "346", "505", "820", "1,100", "数据中心占比持续提升"],
              ], [1500, 1200, 1200, 1200, 1200, 3060], font_size=8.7)

    add_page_break(doc)

    add_heading(doc, "5. 利润、现金流与股本摊薄", 1)
    add_heading(doc, "基准损益预测", 2)
    add_table(doc,
              ["亿美元/美元", "2025A", "2026E", "2027E", "2028E"],
              [
                  ["收入", "346.4", "505.0", "820.0", "1,100.0"],
                  ["Non-GAAP毛利润", "181.7", "281.8", "455.1", "618.2"],
                  ["Non-GAAP毛利率", "52.4%", "55.8%", "55.5%", "56.2%"],
                  ["Non-GAAP经营利润", "77.7", "144.0", "280.0", "418.0"],
                  ["Non-GAAP经营利润率", "22.4%", "28.5%", "34.1%", "38.0%"],
                  ["Non-GAAP净利润", "68.3", "126.2", "242.5", "363.6"],
                  ["摊薄股数（亿股）", "16.4", "16.6", "17.2", "18.0"],
                  ["Non-GAAP EPS", "4.17", "7.60", "14.10", "20.20"],
              ], [3100, 1565, 1565, 1565, 1565], font_size=8.8)
    add_source(doc, "注：2026E采用约13%的Non-GAAP税率；2027-28包含阶段性权证摊薄。由于归属条件复杂，实际股数可能显著偏离。")

    add_heading(doc, "毛利率的核心会计分叉", 2)
    add_text(doc, "MI450与Helios规模化会提升数据中心占比，但不保证毛利率同步上升。若AMD确认更多机架级硬件收入，收入会更高、毛利率可能更低；若以芯片收入为主，收入较低但毛利率更好。本模型将2027毛利率控制在55.5%，没有同时假设“完整机架收入”和“纯芯片毛利率”。")

    add_heading(doc, "权证与战略投资", 2)
    add_table(doc,
              ["项目", "最大规模", "本模型处理", "投资者应关注"],
              [
                  ["OpenAI权证", "1.6亿股", "随采购/技术/股价里程碑分期", "收入增长是否覆盖每股摊薄"],
                  ["Meta权证", "1.6亿股", "随出货与商业条件分期", "定制产品毛利率和集中度"],
                  ["Anthropic投资", "最高50亿美元", "不抵减收入；在估值中作为潜在现金用途", "投资时点、估值和退出路径"],
              ], [1800, 1600, 3100, 2860], font_size=8.5)
    add_callout(doc, "经济摊薄上限", "两项权证合计最高3.2亿股，约相当于当前基础股数的19.6%。即使会计EPS短期尚未体现，估值也不能忽略其条件性经济成本。", fill="FCEBED", accent=RED)

    add_heading(doc, "现金流判断", 2)
    add_text(doc, "AMD为无晶圆厂模式，资本开支强度较低，盈利增长可转化为较强自由现金流。但系统业务放量会提高存货、应收和客户融资需求；同时50亿美元战略投资、NeoCloud支持及潜在回购会限制现金余额的机械累积。")

    add_page_break(doc)

    # Software
    add_heading(doc, "6. 软件生态：迁移成本下降，不等于CUDA壁垒消失", 1)
    add_text(doc, "ROCm.AI与Coding Agent的价值是真实的：它们可帮助代码转换、算子适配和性能调优，降低首次迁移的人力成本。但生产系统的选择由全生命周期TCO决定，不能只看示范代码是否运行。")
    add_table(doc,
              ["层面", "ROCm改善方向", "尚待验证"],
              [
                  ["代码迁移", "Agent自动修改与适配", "复杂自定义算子与长期维护成本"],
                  ["性能", "编译器、内核和框架优化", "不同模型、批次和集群规模下的稳定领先"],
                  ["多机扩展", "通信库和网络协同", "故障恢复、尾延迟与大规模利用率"],
                  ["开发运维", "开放工具链与社区", "调试、分析、版本兼容和企业支持"],
                  ["人才生态", "开放标准降低绑定", "CUDA人才与存量代码的沉没成本"],
              ], [1800, 3300, 4260], font_size=8.8)
    add_heading(doc, "应使用的验证指标", 2)
    add_text(doc, "建议每季跟踪四类硬指标：生产客户数量与规模、集群利用率和稳定性、主流模型性能/美元、从CUDA迁移到ROCm的实际工程周期。只有这些指标连续改善，软件生态才会转化为可持续份额。")
    add_callout(doc, "研究结论", "把ROCm视为份额扩张的必要条件，而不是充分条件。客户合同证明AMD已进入采购池，生产效率决定其能否留在核心池。", fill=LIGHT, accent=BLUE)

    add_page_break(doc)

    # Valuation
    add_heading(doc, "7. 估值与情景分析", 1)
    doc.add_picture(str(chart_value), width=Inches(6.45))
    add_source(doc, "资料来源：本报告测算。目标价值未额外加入净现金溢价，以避免对潜在投资和营运资金形成双重乐观。")

    add_table(doc,
              ["情景", "2027收入", "2027 EPS", "目标PE", "目标价值", "核心条件"],
              [
                  ["审慎", "580亿美元", "9.0美元", "28x", "252美元", "MI450延后、CPU增速回落、摊薄较高"],
                  ["基准", "820亿美元", "14.1美元", "33x", "465美元", "首批GW分阶段交付、EPYC继续扩份额"],
                  ["乐观", "1,100亿美元", "18.0美元", "36x", "648美元", "多客户按期部署、机架交付和毛利率均兑现"],
              ], [1250, 1500, 1350, 1150, 1350, 2760], font_size=8.5)
    add_callout(doc, "评级与目标", "中性 / 观察；基准合理价值465美元，合理区间430-530美元。相较当前484.64美元缺乏足够安全边际，等待Q2业绩、MI450首批收入与摊薄细节验证。", fill="FFF7E6", accent=GOLD)

    add_heading(doc, "为什么不使用640美元作为基准", 2)
    add_text(doc, "640美元基本等于18.28美元2027E EPS乘35倍。其数学没有问题，但把一个接近乐观情景的盈利预测包装为基准情景。若采用14.1美元基准EPS，即使给35倍，价值也只有约494美元。目标价争议的核心是盈利兑现，而不是2-3倍PE差异。")

    add_heading(doc, "估值上修条件", 2)
    for title, body in [
        ("收入条件：", "2026Q3-Q4数据中心收入持续高于基准，且公司披露2027可交付订单。"),
        ("利润条件：", "系统收入放量后Non-GAAP毛利率仍稳定在55%以上。"),
        ("每股条件：", "权证归属速度显著慢于利润增长，战略投资不侵蚀自由现金流。"),
        ("生态条件：", "ROCm生产客户和大规模集群利用率出现可量化改善。"),
    ]:
        add_rich_para(doc, [(title, {"size": 10.3, "bold": True, "color": BLUE}),
                            (body, {"size": 10.3, "bold": False, "color": INK})], after=6)

    add_page_break(doc)

    # Risks and catalysts
    add_heading(doc, "8. 催化剂与风险", 1)
    add_heading(doc, "未来12个月催化剂", 2)
    add_table(doc,
              ["时间窗口", "事件", "正向观察点", "负向观察点"],
              [
                  ["2026Q2财报", "收入、毛利率及Q3指引", "数据中心继续超预期", "H2指引无法支撑爬坡"],
                  ["2026Q3-Q4", "MI450/Helios首批出货", "客户验收和收入确认顺利", "出货与收入确认脱节"],
                  ["2027H1", "Anthropic首个GW开始部署", "多客户并行交付", "融资、供电或集群建设延后"],
                  ["2027", "MI500/Helios 500", "年度迭代、性能和TCO领先", "路线图延误或软件适配滞后"],
              ], [1600, 2400, 2750, 2610], font_size=8.5)

    add_heading(doc, "主要风险排序", 2)
    risks = [
        ("高｜交付与确认风险", "GW合作为多年期框架，系统建设、供电、融资、客户验收均可能推迟收入。"),
        ("高｜客户与权证集中", "少数客户决定大部分增量，同时可能触发显著股本摊薄。"),
        ("高｜软件与利用率", "峰值性能并不等于生产TCO；ROCm稳定性不足会影响复购。"),
        ("中高｜毛利率", "完整机架收入可能稀释毛利率，定制产品也可能让利。"),
        ("中高｜供应链", "HBM、先进封装、光/电互联、液冷与机架集成任何一环均可成为瓶颈。"),
        ("中｜竞争", "英伟达年度迭代、定制ASIC、Arm服务器CPU与英特尔路线图均可能压缩份额。"),
        ("中｜政策", "出口管制、许可与地缘摩擦可能限制产品规格和市场范围。"),
    ]
    for title, body in risks:
        color = RED if title.startswith("高") else GOLD
        add_rich_para(doc, [(title, {"size": 10.2, "bold": True, "color": color}),
                            ("  " + body, {"size": 10.2, "bold": False, "color": INK})], after=7)

    add_heading(doc, "投委会检查清单", 2)
    add_table(doc,
              ["必须回答的问题", "当前状态"],
              [
                  ["每个客户的GPU数量、机架价值量和季度确认计划是否可追踪？", "未充分披露"],
                  ["完整系统与芯片收入的毛利率边界是否清楚？", "需Q3-Q4验证"],
                  ["权证归属和摊薄是否已进入每股模型？", "本报告已做情景处理"],
                  ["50亿美元Anthropic投资及其他融资支持是否进入现金流？", "需公司进一步披露"],
                  ["ROCm是否有生产规模和利用率证据？", "方向改善，数据仍不足"],
              ], [6960, 2400], font_size=8.8)

    add_page_break(doc)

    # Sources and disclaimer
    add_heading(doc, "附录：数据口径、来源与免责声明", 1)
    add_heading(doc, "主要数据口径", 2)
    add_text(doc, "本报告以AMD 2025财年实际数据、2026Q1实际数据和2026Q2公司指引为历史锚。2026-2028年均为研究情景，不是公司指引。Non-GAAP指标用于跨期经营比较；估值同时考虑潜在权证摊薄，不把最高权证股数机械视为必然归属。")
    add_heading(doc, "主要公开来源", 2)
    sources = [
        "1. AMD Reports First Quarter 2026 Financial Results，2026-05-05：Q1业绩与Q2指引。",
        "2. AMD and OpenAI Announce Strategic Partnership to Deploy 6 Gigawatts of AMD GPUs，2025-10-06。",
        "3. AMD and Meta Announce Expanded Strategic Partnership to Deploy 6 Gigawatts of AMD GPUs，2026-02-24。",
        "4. AMD and Anthropic Announce Strategic Partnership to Deploy Up to 2 Gigawatts，2026-07-22。",
        "5. AMD Advancing AI 2026 Keynote及产品路线图，2026-07-23。",
        "6. AMD 2025 Form 10-K与2026Q1财务报告。",
        "7. UBS, Revisiting AMD and ARM Estimates Amid Agentic AI Acceleration，2026-06-24。",
        "8. 华泰证券，《软硬生态全面进击，剑指Agentic AI算力变局》，2026-07-26。",
        "9. 公开市场价格数据，截至2026-08-04 00:15 UTC。",
    ]
    for s in sources:
        add_text(doc, s, size=9.4, color=INK, after=5)
    add_heading(doc, "关键网址", 2)
    urls = [
        "https://www.amd.com/en/newsroom/press-releases/2026-5-5-amd-reports-first-quarter-2026-financial-results.html",
        "https://www.amd.com/en/newsroom/press-releases/2025-10-6-amd-and-openai-announce-strategic-partnership-to-d.html",
        "https://www.amd.com/en/newsroom/press-releases/2026-2-24-amd-and-meta-announce-expanded-strategic-partnersh.html",
        "https://ir.amd.com/news-events/press-releases/detail/1292/amd-and-anthropic-announce-strategic-partnership-to-deploy-up-to-2-gigawatts-of-amd-instinct-mi450-series-gpus",
        "https://www.amd.com/en/corporate/events/advancing-ai.html",
    ]
    for u in urls:
        add_text(doc, u, size=8.2, color=BLUE, after=4)
    add_heading(doc, "免责声明", 2)
    add_text(doc, "本报告仅用于研究讨论，不构成任何证券买卖建议、要约或保证。预测高度依赖产品交付、客户部署、会计确认、竞争和市场条件，实际结果可能显著不同。报告使用的公开信息被认为可靠，但不保证完整或无误。投资者应结合自身风险承受能力独立判断。", size=9.2, color=MID)

    # Document properties
    props = doc.core_properties
    props.title = "AMD 2026Q2 Codex独立研报：AI算力第二供应链的真实弹性与兑现门槛"
    props.subject = "AMD公司研究与三情景估值"
    props.author = "Independent Research Draft"
    props.keywords = "AMD, MI450, EPYC, ROCm, Helios, valuation"
    props.comments = "As of 2026-08-04, pre-Q2 earnings"

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
